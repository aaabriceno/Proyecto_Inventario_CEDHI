"""Agrega una columna "MODULO" vacia al FINAL de cada hoja con datos de los
Excel de datos_iniciales/INVENTARIO_2026/.

No llena ningun valor: el CEDHI/usuario debe completar manualmente que
modulo corresponde a cada fila, ya que el Excel original no trae esa
informacion y el sistema no puede inferirla por si solo.

La columna se agrega al final (no insertada junto a UBICACION) porque estas
hojas usan celdas combinadas (merged cells) alrededor de UBICACION/FECHA/
ESTADO: insertar una columna en medio de un merge lo corrompe (confirmado
en una copia de prueba -- el header "ESTADO" y sus sub-columnas B/R/M
quedaban mal desplazados). Agregar al final no toca ningun merge existente.

Uso: python3 scripts/agregar_columna_modulo.py
(correr desde la raiz del repo inventario_cedhi, con openpyxl instalado
-- usar el venv del bench: source ../../env/bin/activate)
"""

import os

import openpyxl

FOLDER = os.path.join(os.path.dirname(__file__), "..", "datos_iniciales", "INVENTARIO_2026")


def encontrar_fila_ubicacion(ws):
    for row_idx in range(1, 15):
        for col_idx in range(1, ws.max_column + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            if value and "UBICAC" in str(value).upper():
                return row_idx
    return None


def main():
    procesados = []
    for fname in sorted(os.listdir(FOLDER)):
        if not fname.lower().endswith(".xlsx") or fname.startswith("Copia"):
            continue
        path = os.path.join(FOLDER, fname)
        wb = openpyxl.load_workbook(path)
        cambiado = False

        for sheet in wb.sheetnames:
            ws = wb[sheet]
            header_row = encontrar_fila_ubicacion(ws)
            if header_row is None:
                continue

            modulo_col = ws.max_column + 1
            existing = ws.cell(row=header_row, column=modulo_col).value
            if existing and "MODULO" in str(existing).upper():
                continue

            ws.cell(row=header_row, column=modulo_col, value="MODULO")
            cambiado = True
            procesados.append(f"{fname} | {sheet}")

        if cambiado:
            wb.save(path)

    print(f"Hojas modificadas: {len(procesados)}")
    for p in procesados:
        print(" -", p)


if __name__ == "__main__":
    main()
